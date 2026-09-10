import os,time,io,json
os.environ['APP_MODE']='demo'
os.environ['DATABASE_URL']='sqlite:///./test-workflow.db'
os.environ['API_KEY']='test-key'
from fastapi.testclient import TestClient
from app.main import app
from app.demo import KADAPA
def test_discovery_and_exports_match():
    with TestClient(app,headers={'X-API-Key':'test-key'}) as c:
        suggestions=c.get('/api/locations/autocomplete',params={'q':'kad','mode':'demo'}).json()
        assert len(suggestions)==2
        r=c.post('/api/search-jobs',json={'terms':['orthopaedics','bone & joint','joint replacement','spine clinic'],'location':KADAPA,'scope':'radius','radius_km':25,'mode':'demo'})
        assert r.status_code==201,r.text
        job=r.json()
        for _ in range(200):
            job=c.get('/api/search-jobs/'+job['id']).json()
            if job['stage'] in ('completed','failed','partially_completed'): break
            time.sleep(.03)
        assert job['stage']=='completed',job
        assert job['raw_observations']==13
        assert job['unique_entities']==12
        result=c.get('/api/entities',params={'job_id':job['id'],'page_size':50}).json()
        assert result['total']==12
        entity=result['items'][0]
        profile=c.get('/api/entities/'+entity['id']).json()
        assert profile['evidence'] and profile['sources']
        assert c.get('/api/entities',params={'job_id':job['id'],'search':'no-such-provider'}).json()['total']==0
        for format in ['json','csv','xlsx']:
            ex=c.post('/api/exports',json={'job_id':job['id'],'format':format,'scope':'selected','entity_ids':[entity['id']]}).json()
            for _ in range(100):
                ex=c.get('/api/exports/'+ex['id']).json()
                if ex['status']!='queued':break
                time.sleep(.02)
            assert ex['status']=='completed',ex
            download=c.get('/api/exports/'+ex['id']+'/download')
            assert download.status_code==200
            if format=='json': assert download.json()['businesses'][0]['id']==entity['id']
            elif format=='xlsx':
                from openpyxl import load_workbook
                wb=load_workbook(io.BytesIO(download.content));assert wb.sheetnames==['Businesses','Categories','Sources','Search Metadata']
                assert wb['Businesses'].max_row==2
        assert c.post('/api/search-jobs',json={'terms':[]}).status_code==422
