import csv,io,json,time
from pathlib import Path
from sqlalchemy import select
import pandas as pd
from openpyxl.styles import Font,PatternFill
from .config import settings
from .models import SearchJob,ExportJob
from .repository import query_entities,ordered,serialize_entity
from .schemas import ExportRequest
from .db import Session
def safe_cell(value):
    if isinstance(value,(dict,list)): value=json.dumps(value,ensure_ascii=False)
    if isinstance(value,str) and value.lstrip().startswith(('=','+','-','@','\t','\r')): return "'"+value
    return value
def export_bytes(rows,metadata,format):
    flat=[{k:v for k,v in r.items() if k not in ('evidence','sources','categories')} for r in rows]
    if format=='json': return json.dumps({'businesses':rows,'search_metadata':metadata},ensure_ascii=False,indent=2).encode()
    if format=='csv':
        stream=io.StringIO(newline='')
        fields=list(flat[0]) if flat else ['id','name']
        writer=csv.DictWriter(stream,fieldnames=fields);writer.writeheader()
        writer.writerows([{k:safe_cell(v) for k,v in r.items()} for r in flat])
        return stream.getvalue().encode('utf-8-sig')
    output=io.BytesIO()
    categories=[{'entity_id':r['id'],'dimension':dim,'value':v} for r in rows for dim in ['categories','specialties','subspecialties','services'] for v in r.get(dim,[])]
    sources=[{'entity_id':r['id'],**s} for r in rows for s in r.get('sources',[])]
    sheets={'Businesses':flat,'Categories':categories,'Sources':sources,'Search Metadata':[{'key':k,'value':v} for k,v in metadata.items()]}
    with pd.ExcelWriter(output,engine='openpyxl') as writer:
        for name,data in sheets.items():
            frame=pd.DataFrame([{k:safe_cell(v) for k,v in r.items()} for r in data])
            if frame.empty: frame=pd.DataFrame(columns=['No records'])
            frame.to_excel(writer,sheet_name=name,index=False)
            ws=writer.sheets[name];ws.freeze_panes='A2';ws.auto_filter.ref=ws.dimensions
            for cell in ws[1]: cell.font=Font(bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='156E60')
            for col in ws.columns:
                ws.column_dimensions[col[0].column_letter].width=min(60,max(16,max(len(str(c.value or '')) for c in col[:100])+2))
    return output.getvalue()
def run_export(export_id,session_factory=Session):
    started=time.monotonic()
    with session_factory() as session:
        export=session.get(ExportJob,export_id)
        if not export or export.status!='queued': return
        try:
            request=ExportRequest.model_validate(export.request);job=session.get(SearchJob,request.job_id)
            q=query_entities(job,request.filters if request.scope=='filtered' and session.bind.dialect.name!='sqlite' else {})
            if request.scope=='selected':
                from .models import BusinessEntity
                q=q.where(BusinessEntity.id.in_(request.entity_ids))
            
            if session.bind.dialect.name=='sqlite' and request.scope=='filtered':
                from .main import filtered_rows
                from .models import BusinessEntity
                from uuid import UUID
                ids=[UUID(r['id']) for r in filtered_rows(session,job,request.filters,'name')]
                q=query_entities(job,{}).where(BusinessEntity.id.in_(ids))
            entities=session.scalars(ordered(q,'name',job)).all()
            rows=[serialize_entity(session,e,job,True) for e in entities]
            metadata={**job.request,'normalized_terms':job.normalized_terms,'generated_query_variants':job.query_variants,'job_id':str(job.id),'run_date':str(job.created_at),'raw_observation_count':job.raw_observations,'unique_entity_count':job.unique_entities,'exported_entity_count':len(rows),'source_adapters':sorted({s['source'] for r in rows for s in r['sources']})}
            payload=export_bytes(rows,metadata,request.format)
            export.payload=payload;export.status='completed'
        except Exception as exc: export.status='failed';export.error=str(exc)[:500]
        export.duration_seconds=round(time.monotonic()-started,2);session.commit()
